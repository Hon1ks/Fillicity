"""Turn user-supplied files (md, txt, xlsx, docx, pdf, images) into SourceDocuments
that the AI engine can reason over.
"""
from __future__ import annotations

import base64
import os

from app.core.models import SourceDocument

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}
MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".gif": "image/gif",
}

SUPPORTED_EXTS = IMAGE_EXTS | {".md", ".txt", ".csv", ".xlsx", ".xlsm", ".docx", ".pdf"}


class UnsupportedFileError(Exception):
    pass


def load_source(path: str) -> SourceDocument:
    ext = os.path.splitext(path)[1].lower()
    name = os.path.basename(path)

    if ext not in SUPPORTED_EXTS:
        raise UnsupportedFileError(f"Формат не поддерживается: {ext or '(без расширения)'}")

    if ext in IMAGE_EXTS:
        with open(path, "rb") as f:
            payload = base64.standard_b64encode(f.read()).decode("ascii")
        return SourceDocument(
            path=path, name=name, kind="image",
            image_b64=payload, media_type=MEDIA_TYPES[ext],
        )

    if ext in (".md", ".txt", ".csv"):
        text = _read_plain_text(path)
    elif ext in (".xlsx", ".xlsm"):
        text = _read_xlsx(path)
    elif ext == ".docx":
        text = _read_docx(path)
    elif ext == ".pdf":
        text = _read_pdf(path)
    else:  # pragma: no cover - guarded by SUPPORTED_EXTS above
        raise UnsupportedFileError(ext)

    return SourceDocument(path=path, name=name, kind="text", text=text)


def _read_plain_text(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _read_xlsx(path: str) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"# Лист: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _read_docx(path: str) -> str:
    import docx

    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pdf(path: str) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    return "\n".join(parts)
